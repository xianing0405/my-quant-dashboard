#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""内部知识库服务层测试。

本文件只使用「明确标注的测试资料」，且全部写入临时目录，绝不进入正式
raw_docs / store，因此不会污染正式知识库。

验证分层（见底部）：
- 纯服务层测试（无模型依赖，确定性通过）
- mock 模型测试（monkeypatch llm，验证接口与异常路径）
- 真实模型 / 真实检索验证：本文件不运行（缺凭证），列为「未验证」
"""

from __future__ import annotations

import json
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
WEBSITE = os.path.dirname(HERE)
if WEBSITE not in sys.path:
    sys.path.insert(0, WEBSITE)

from core import knowledge as kb  # noqa: E402
from core.knowledge import storage  # noqa: E402
from core.knowledge.auth import KnowledgeAuthError  # noqa: E402
from core.knowledge.models import PARSE_NEEDS_OCR, PARSE_READY  # noqa: E402

USER = "admin"


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("KNOWLEDGE_ACCESS_MODE", "open")
    monkeypatch.setenv("KNOWLEDGE_STORE_DIR", str(tmp_path / "store"))
    monkeypatch.setenv("KNOWLEDGE_MATERIALS_DIR", str(tmp_path / "materials"))
    os.makedirs(tmp_path / "materials", exist_ok=True)
    storage.reset_store()
    yield tmp_path
    storage.reset_store()


# ---------------------------------------------------------------------------
# 测试夹具生成（独立测试目录，临时）
# ---------------------------------------------------------------------------
def _write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def _make_pdf(path, pages):
    """用 matplotlib 生成带文本层的多页 PDF（ASCII，避免 CJK 字体缺失）。"""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.backends.backend_pdf import PdfPages
    import matplotlib.pyplot as plt
    with PdfPages(str(path)) as pdf:
        for i, txt in enumerate(pages, 1):
            fig = plt.figure()
            fig.text(0.1, 0.5, txt)
            pdf.savefig(fig)
            plt.close(fig)


def _make_docx(path, heading, paras, table=None):
    from docx import Document
    d = Document()
    d.add_heading(heading, level=1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    d.save(str(path))


# ---------------------------------------------------------------------------
# 1. 从 PDF 找到指定观点并定位正确页码
# ---------------------------------------------------------------------------
def test_pdf_page_number(env):
    pdf = env / "materials" / "研报_20260910.pdf"
    _make_pdf(pdf, [
        "Page one: SiC demand is rising.",
        "Page two: team view: SiC valuation is cheap.",
    ])
    r = kb.ingest_file(str(pdf), user=USER, source_category="外部研报", publish_date="2026-09-10")
    assert r["status"] == PARSE_READY

    res = kb.search_materials("SiC valuation cheap", USER, top_k=5)
    assert res["hits"], "应命中 PDF 片段"
    # 该观点在第二页
    hit = next(h for h in res["hits"] if "valuation is cheap" in h["text"].lower())
    assert hit["page_number"] == 2
    assert hit["title"].startswith("研报_20260910")


# ---------------------------------------------------------------------------
# 2. 从对话区分提问、条件性观点和回答
# ---------------------------------------------------------------------------
def test_dialogue_roles(env):
    d = env / "materials" / "对话_20260916.json"
    _write(d, json.dumps({
        "format": "dialogue",
        "messages": [
            {"speaker": "张三", "role": "question", "time": "2026-09-16T14:30:00", "text": "碳化硅现在怎么看？"},
            {"speaker": "李四", "role": "answer", "time": "2026-09-16T14:32:00", "text": "若美联储降息，则关注碳化硅。"},
            {"speaker": "王五", "role": "statement", "time": "2026-09-16T14:35:00", "text": "但需警惕产能过剩风险。"},
        ],
    }, ensure_ascii=False))
    r = kb.ingest_file(str(d), user=USER, source_category="对话", publish_date="2026-09-16")
    assert r["status"] == PARSE_READY

    detail = kb.document_evidence(r["document_id"], USER)
    by_speaker = {c["speaker"]: c for c in detail["chunks"]}
    assert by_speaker["张三"]["role"] == "question"
    # 条件性观点保留条件「若」
    assert "若美联储降息" in by_speaker["李四"]["text"]
    # 回答关联到前面的提问（连续问答关系）
    assert by_speaker["李四"]["parent_evidence_id"] == by_speaker["张三"]["evidence_id"]


# ---------------------------------------------------------------------------
# 3. 早期观点与后续修订，不混淆时间（历史可追溯）
# ---------------------------------------------------------------------------
def test_viewpoint_revision_history(env):
    vid1 = kb.add_viewpoint(USER, research_subject="碳化硅",
                            core_judgment="若降息则关注碳化硅", viewpoint_date="2026-09-10",
                            proposer="张三")
    kb.set_viewpoint_status(vid1, USER, "已修订")
    hist = kb.get_viewpoint_history("碳化硅", USER)
    assert len(hist["viewpoints"]) == 2, "修订应产生新记录并保留旧记录"
    dates = [v["viewpoint_date"] for v in hist["viewpoints"]]
    assert dates == sorted(dates), "应按时间正序"
    # 新记录通过 relationships 指向旧观点
    new_vp = next(v for v in hist["viewpoints"] if v["viewpoint_id"] != vid1)
    assert new_vp["relationships"][0]["target_viewpoint_id"] == vid1


# ---------------------------------------------------------------------------
# 4. 无证据时不编造答案或引用
# ---------------------------------------------------------------------------
def test_no_evidence_no_fabrication(env):
    res = kb.search_materials("完全不存在的主题 XYZ123", USER, top_k=3)
    assert res["hits"] == []
    # 页面层 _kb_generate_answer 的行为：无命中返回明确无依据文案
    from web_dashboard import _kb_generate_answer
    answer, hits, cited = _kb_generate_answer("完全不存在的主题 XYZ123", [])
    assert "未找到支持材料" in answer
    assert hits == [] and cited == []


# ---------------------------------------------------------------------------
# 5. 日期筛选不引入截止日期之后的材料
# ---------------------------------------------------------------------------
def test_date_cutoff(env):
    a = env / "materials" / "A_20260910.txt"
    b = env / "materials" / "B_20260925.txt"
    _write(a, "碳化硅 需求上行")
    _write(b, "碳化硅 需求上行 最新数据")
    kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")
    kb.ingest_file(str(b), user=USER, source_category="内部观点", publish_date="2026-09-25")

    res = kb.search_materials("碳化硅", USER, top_k=10, filters={"date_to": "2026-09-20"})
    titles = {h["title"] for h in res["hits"]}
    assert "B_20260925" not in titles, "截止日期后的材料不应进入"
    assert "A_20260910" in titles


def test_date_unknown_flagged(env):
    a = env / "materials" / "nodate.txt"
    _write(a, "碳化硅 需求")
    kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date=None)
    res = kb.search_materials("碳化硅", USER, top_k=10, filters={"date_to": "2026-09-20"})
    assert res["date_unknown_excluded"] == 1, "日期未知材料应单独计数而非悄悄纳入"


# ---------------------------------------------------------------------------
# 6. 删除资料后不能继续从索引召回
# ---------------------------------------------------------------------------
def test_delete_removes_from_index(env):
    a = env / "materials" / "del.txt"
    _write(a, "光纤 景气度上行")
    r = kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")
    assert kb.search_materials("光纤", USER)["hits"]
    kb.delete_document(r["document_id"], user=USER)
    assert kb.search_materials("光纤", USER)["hits"] == [], "删除后不应再召回"


# ---------------------------------------------------------------------------
# 7. 无权限用户不能搜索或读取内部资料
# ---------------------------------------------------------------------------
def test_auth_denies_anonymous(env, monkeypatch):
    a = env / "materials" / "secret.txt"
    _write(a, "内部敏感观点")
    kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")

    monkeypatch.setenv("KNOWLEDGE_ACCESS_MODE", "disabled")
    with pytest.raises(KnowledgeAuthError):
        kb.search_materials("敏感", user=None)
    with pytest.raises(KnowledgeAuthError):
        kb.list_documents(user=None)


# ---------------------------------------------------------------------------
# 8. 重复上传不会重复入库
# ---------------------------------------------------------------------------
def test_duplicate_upload_dedup(env):
    a = env / "materials" / "dup.txt"
    _write(a, "重复内容测试")
    r1 = kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")
    r2 = kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")
    assert r1["status"] == PARSE_READY
    assert r2["status"] == "duplicate"
    assert len(kb.list_documents(USER)) == 1


# ---------------------------------------------------------------------------
# 9. 新点评能够列出支持与冲突证据（mock 模型）
# ---------------------------------------------------------------------------
def test_commentary_lists_evidence(env, monkeypatch):
    sup = env / "materials" / "支持_20260910.txt"
    _write(sup, "碳化硅：需求上行，若降息则受益。")
    kb.ingest_file(str(sup), user=USER, source_category="内部观点", publish_date="2026-09-10")

    captured = {}
    from core.knowledge import llm as kb_llm

    def _fake_complete(prompt, max_tokens=2000, system=None):
        captured["prompt"] = prompt
        return "## 三、支持 / 削弱 / 尚不能验证的证据\n- 支持：[证据:1] 需求上行\n- 削弱：产能过剩\n## 四、本次分析与不确定性\n- 尚需验证订单"

    monkeypatch.setattr(kb_llm, "complete", _fake_complete)
    monkeypatch.setattr(kb_llm, "available", lambda: True)

    r = kb.build_commentary("某公司 2026-09-28 公告碳化硅扩产", USER, subject="碳化硅")
    assert r["ok"] is True
    assert r["evidence_count"] >= 1
    assert r["ai_generated"] is True
    # 用户粘贴材料标注来源未核验
    assert "用户提供，来源未独立核验" in r["output_markdown"]
    # 检索到的证据编号进入提示词，供模型引用
    assert "[证据:1]" in captured["prompt"]


# ---------------------------------------------------------------------------
# 10. 模型不可用 / 解析失败时显示准确状态
# ---------------------------------------------------------------------------
def test_parse_failure_status(env):
    bad = env / "materials" / "bad.bin"
    bad.write_bytes(b"\x00\x01\x02")
    # 不支持的扩展名：ingest_file 会失败（解析失败）
    r = kb.ingest_file(str(bad), user=USER, source_category="其他")
    assert r["status"] in ("解析失败", "需要OCR")


def test_scanned_pdf_needs_ocr(env):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    pdf = env / "materials" / "scan.pdf"
    fig = plt.figure()  # 无文本
    fig.savefig(str(pdf))
    plt.close(fig)
    r = kb.ingest_file(str(pdf), user=USER, source_category="外部研报")
    assert r["status"] == PARSE_NEEDS_OCR, "无文本层 PDF 应标记需要 OCR，而非成功"


def test_llm_unavailable_status(env, monkeypatch):
    from core.knowledge import llm as kb_llm
    monkeypatch.setattr(kb_llm, "available", lambda: False)
    a = env / "materials" / "x.txt"
    _write(a, "碳化硅 需求")
    r = kb.ingest_file(str(a), user=USER, source_category="内部观点", publish_date="2026-09-10")
    ext = kb.extract_viewpoints_from_document(r["document_id"], USER)
    assert ext["ok"] is False and "模型不可用" in ext["reason"]


# ---------------------------------------------------------------------------
# 验证分层说明（供人工与 CI 区分）
# ---------------------------------------------------------------------------
def test_verification_tiers():
    """本函数仅作说明，恒通过：记录三类验证状态。"""
    print("\n验证分层：")
    print("  - mock 测试通过：commentary / 观点提取（monkeypatch llm）")
    print("  - 真实检索已验证：关键词检索、页码定位、去重、删除、权限、日期筛选（无模型依赖）")
    print("  - 尚因缺少配置而无法验证：真实大模型问答 / 观点 AI 提取 / 语义检索（需 ANTHROPIC_AUTH_TOKEN 与 EMBEDDING_*）")
    assert True
